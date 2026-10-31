import io
import shutil
import subprocess
from pathlib import Path

from django.conf import settings


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".doc",
    ".xlsx",
    ".xls",
    ".csv",
    ".txt",
    ".jpg",
    ".jpeg",
    ".png",
    ".webp",
}


def extract_text(uploaded_file):
    """
    Extract searchable text from common company document formats.

    The function accepts normal Django uploaded files as well as
    file-like objects retrieved from Cloudinary RAW storage.

    Returns:

    {
        "text": "...",
        "method": "docx",
        "error": "",
    }
    """

    filename = getattr(uploaded_file, "name", "") or ""

    extension = Path(filename).suffix.lower()

    try:
        if extension == ".docx":
            return _extract_docx(uploaded_file)

        if extension == ".doc":
            return {
                "text": "",
                "method": "unsupported",
                "error": (
                    "Legacy .doc extraction is not currently supported. "
                    "Please convert the document to .docx or PDF."
                ),
            }

        if extension == ".pdf":
            return _extract_pdf(uploaded_file)

        if extension in {".jpg", ".jpeg", ".png", ".webp"}:
            return _extract_image_ocr(uploaded_file)

        if extension in {".xlsx", ".xls"}:
            return _extract_excel(uploaded_file)

        if extension == ".csv":
            return _extract_csv(uploaded_file)

        if extension == ".txt":
            return _extract_text_file(uploaded_file)

        return {
            "text": "",
            "method": "unsupported",
            "error": (
                f"Content extraction is not currently supported "
                f"for {extension or 'this file type'}."
            ),
        }

    except Exception as exc:
        return {
            "text": "",
            "method": "error",
            "error": str(exc),
        }


def _extract_docx(uploaded_file):
    from docx import Document

    uploaded_file.seek(0)

    document = Document(uploaded_file)

    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            values = []

            for cell in row.cells:
                value = cell.text.strip()

                if value:
                    values.append(value)

            if values:
                parts.append(" | ".join(values))

    return {
        "text": "\n".join(parts),
        "method": "docx",
        "error": "",
    }


def _extract_pdf(uploaded_file):
    import fitz

    uploaded_file.seek(0)

    pdf = fitz.open(
        stream=uploaded_file.read(),
        filetype="pdf",
    )

    parts = []
    ocr_page_count = 0
    max_ocr_pages = getattr(settings, "COMPANY_DOCUMENT_OCR_MAX_PAGES", 20)
    limit_reached = False

    try:
        for page in pdf:
            text = page.get_text("text").strip()
            if text:
                parts.append(text)
                continue
            if ocr_page_count >= max_ocr_pages:
                limit_reached = True
                continue
            image_bytes = page.get_pixmap(
                matrix=fitz.Matrix(2, 2),
                alpha=False,
            ).tobytes("png")
            parts.append(_run_tesseract(image_bytes))
            ocr_page_count += 1
    finally:
        pdf.close()
    used_ocr = ocr_page_count > 0
    return {
        "text": "\n".join(part for part in parts if part),
        "method": "pdf_ocr" if used_ocr else "pdf",
        "error": (
            f"OCR was limited to {max_ocr_pages} scanned pages; remaining scanned "
            "pages were not processed."
            if limit_reached
            else ""
        ),
    }


def _extract_image_ocr(uploaded_file):
    from PIL import Image, UnidentifiedImageError

    uploaded_file.seek(0)
    try:
        with Image.open(uploaded_file) as image:
            if image.width * image.height > 50_000_000:
                raise ValueError("Image dimensions exceed the OCR processing limit.")
            image.load()
            normalized = image.convert("RGB")
    except (UnidentifiedImageError, OSError) as exc:
        raise ValueError("The uploaded image could not be decoded.") from exc

    buffer = io.BytesIO()
    normalized.save(buffer, format="PNG")
    return {
        "text": _run_tesseract(buffer.getvalue()),
        "method": "image_ocr",
        "error": "",
    }


def _run_tesseract(image_bytes):
    executable = getattr(settings, "COMPANY_DOCUMENT_TESSERACT_CMD", "tesseract")
    resolved_executable = shutil.which(executable)
    if not resolved_executable:
        raise RuntimeError(
            "Local OCR is unavailable. Install Tesseract or configure "
            "COMPANY_DOCUMENT_TESSERACT_CMD."
        )

    result = subprocess.run(
        [
            resolved_executable,
            "stdin",
            "stdout",
            "-l",
            getattr(settings, "COMPANY_DOCUMENT_OCR_LANGUAGE", "eng"),
        ],
        input=image_bytes,
        capture_output=True,
        timeout=20,
        check=False,
    )
    if result.returncode:
        message = result.stderr.decode("utf-8", errors="replace").strip()
        raise RuntimeError(
            f"Tesseract OCR failed: {message[:500] or 'unknown OCR error'}"
        )
    return result.stdout.decode("utf-8", errors="replace").strip()


def _extract_excel(uploaded_file):
    from openpyxl import load_workbook

    uploaded_file.seek(0)

    workbook = load_workbook(
        uploaded_file,
        read_only=True,
        data_only=True,
    )

    parts = []

    for worksheet in workbook.worksheets:
        parts.append(
            f"WORKSHEET: {worksheet.title}"
        )

        for row in worksheet.iter_rows(
            values_only=True
        ):
            values = [
                str(value).strip()
                for value in row
                if value is not None
                and str(value).strip()
            ]

            if values:
                parts.append(" | ".join(values))

    workbook.close()

    return {
        "text": "\n".join(parts),
        "method": "excel",
        "error": "",
    }


def _extract_csv(uploaded_file):
    import csv
    import io

    uploaded_file.seek(0)

    raw = uploaded_file.read()

    text = raw.decode(
        "utf-8",
        errors="replace",
    )

    reader = csv.reader(
        io.StringIO(text)
    )

    parts = []

    for row in reader:
        values = [
            value.strip()
            for value in row
            if value.strip()
        ]

        if values:
            parts.append(" | ".join(values))

    return {
        "text": "\n".join(parts),
        "method": "csv",
        "error": "",
    }


def _extract_text_file(uploaded_file):
    uploaded_file.seek(0)

    raw = uploaded_file.read()

    if isinstance(raw, bytes):
        text = raw.decode(
            "utf-8",
            errors="replace",
        )
    else:
        text = str(raw)

    return {
        "text": text,
        "method": "text",
        "error": "",
    }